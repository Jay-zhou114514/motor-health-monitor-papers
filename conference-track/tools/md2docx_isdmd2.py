#!/usr/bin/env python3
"""Second pass of the ISDMD docx build: table captions, in-text table references, the display
equation as a Word equation object, the figure placed in the results section, and references
generated from REFERENCES.md in the template's IOP-CS-ReferenceText style."""

from __future__ import annotations

import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

TRACK = r'C:\Users\32597\Documents\Codex\2026-09-15\github\motor-health-monitor-papers\conference-track'
MD = os.path.join(TRACK, 'PAPER_v1.0_submission-draft.md')
REFS = os.path.join(TRACK, 'REFERENCES.md')
TEMPLATE = os.path.join(TRACK, 'latex', 'isdmd-template', 'ISDMD 2026 Template',
                        'WordGuidelines', 'WordGuidelines', 'IOP-ConfSer-template.docx')
OUT = os.path.join(TRACK, 'latex', 'PAPER_v1.0_ISDMD.docx')
FIGURE = os.path.join(TRACK, 'figures', 'fig1_three_faces.png')


def read(path):
    with io.open(path, 'r', encoding='utf-8', newline='') as fh:
        return fh.read()


def strip_markup(text: str) -> str:
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'\1', text)
    return re.sub(r'`([^`]+)`', r'\1', text)


doc = Document(TEMPLATE)
body = doc.element.body
for child in list(body):
    if not child.tag.endswith('}sectPr'):
        body.remove(child)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.0


def para(text='', style=None, align=None, italic=False):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(strip_markup(text))
        run.italic = italic
    if align is not None:
        p.alignment = align
    return p


def equation(text, number):
    """Insert an editable OMML equation (Microsoft's native equation format) with its number."""
    p = doc.add_paragraph()
    omath = OxmlElement('m:oMathPara')
    m = OxmlElement('m:oMath')
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = text
    r.append(t)
    m.append(r)
    omath.append(m)
    p._p.append(omath)
    run = p.add_run('\t\t(' + str(number) + ')')
    run.font.size = Pt(10.5)
    return p


lines = read(MD).split('\n')
title = 'Three Faces of Evaluation Uncertainty in Healthy-Data-Only Bearing Anomaly Detection'
para(title, style='Heading 1', align=WD_ALIGN_PARAGRAPH.CENTER)
para('Author One*, Author Two', align=WD_ALIGN_PARAGRAPH.CENTER)
para('Affiliation, City, Postal Code, Country', align=WD_ALIGN_PARAGRAPH.CENTER)
para('*Corresponding author: name@example.com', align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

abstract, in_abs = [], False
for line in lines:
    s = line.strip()
    if s.startswith('## Abstract'):
        in_abs = True
        continue
    if in_abs:
        if s.startswith('---') or s.startswith('## '):
            break
        if s:
            abstract.append(s)
para('Abstract', style='Heading 2')
para(' '.join(abstract))

table_no = 0
eq_no = 0
inserted_figure = False
pending_caption = None
i = 0
while i < len(lines):
    s = lines[i].strip()

    if s.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i])
            i += 1
        cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
        header, body_rows = cells[0], cells[2:]
        table_no += 1
        table = doc.add_table(rows=1, cols=len(header))
        table.style = 'Table Grid'
        for j, cell in enumerate(header):
            table.rows[0].cells[j].text = strip_markup(cell)
        for row in body_rows:
            cells_row = table.add_row().cells
            for j, cell in enumerate(row[:len(header)]):
                cells_row[j].text = strip_markup(cell)
        para(f'Table {table_no}. {pending_caption or "Summary of results."}', style=None)
        pending_caption = None
        continue

    if s.startswith('## '):
        section = s[3:].strip()
        if section.lower() == 'title':
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('---'):
                i += 1
            continue
        if section != '1. Introduction':
            para(section, style='Heading 1')
        if section.startswith('4.') and not inserted_figure and os.path.isfile(FIGURE):
            doc.add_picture(FIGURE, width=Inches(6.0))
            para('Figure 1. Three faces of evaluation uncertainty in healthy-data-only bearing '
                 'anomaly detection: evaluation scope (a), independent units (b), the '
                 'healthy/degraded boundary (c), and the reporting checklist (d).')
            inserted_figure = True
    elif s.startswith('### '):
        para(s[4:].strip(), style='Heading 2')
    elif s and not s.startswith('>') and not re.match(r'^(\d+\.\s|- |---)', s):
        # display equation in 3.3 becomes a Word equation object
        if 'H = clip(max(20, 0.10·N), 20, 60)' in s and 'swept five variants' in s:
            head = s.split('We froze one healthy-phase rule')[0]
            para(head)
            eq_no += 1
            equation('H = clip(max(20, 0.10·N), 20, 60)', eq_no)
            para('with H/N ≤ 0.25. We swept five variants of it — fraction 5/10/20% and floor '
                 '10/20/30 — on two independently collected run-to-failure datasets.')
        else:
            text = strip_markup(s)
            if 'healthy/degraded boundary number 7, 7 and 3' in text:
                text += ' (Table 3)'
            if 'SD in this table is the pooled standard deviation' in text:
                text += ' (Table 2)'
            if text.startswith('Detector, threshold rule and dataset are held fixed'):
                text += ' (Table 1)'
            para(text)
            for caption in ('Evaluation scope and reported false-alarm rate.',
                            'Unit effect by detector, fixed record budget.',
                            'Healthy-phase boundary sweep.'):
                if pending_caption is None and text.startswith(
                        {'Evaluation scope and reported false-alarm rate.': 'Detector, threshold',
                         'Unit effect by detector, fixed record budget.': 'The sample size is fixed',
                         'Healthy-phase boundary sweep.': 'Run-to-failure datasets carry no onset'}
                        [caption]):
                    pending_caption = caption
    i += 1

# ---------------------------------------------------------------- references
ref_lines = []
for line in read(REFS).split('\n'):
    s = line.strip()
    m = re.match(r'^\[?(\d+)\]?[.)]?\s+(.*\S)', s)
    if m:
        ref_lines.append((m.group(1), strip_markup(m.group(2))))
if ref_lines:
    para('References', style='Heading 1')
    for num, text in ref_lines:
        p = doc.add_paragraph()
        try:
            p.style = doc.styles['IOP-CS-ReferenceText']
        except KeyError:
            pass
        p.add_run(f'[{num}] {text}')
    print(f'references: {len(ref_lines)}')
else:
    print('references: none parsed (REFERENCES.md format not recognised)')

doc.save(OUT)
print('wrote', OUT, f'| tables={table_no} equations={eq_no} figure={inserted_figure}')
